#!/usr/bin/env python3
"""X アカウントの全ポストを時系列Excelにする。使い方: x_account_posts_xlsx.py <@IDまたはプロフィールURL> <出力ディレクトリ>"""
import json, os, re, subprocess, sys, time, urllib.parse
from datetime import datetime, timezone, timedelta

API = "https://api.fxtwitter.com"
JST = timezone(timedelta(hours=9))


def get(url, retries=4):
    d = None
    for t in range(retries):
        o = subprocess.run(["curl", "-s", url], capture_output=True, text=True).stdout
        try:
            d = json.loads(o)
            if d.get("results") or d.get("user") or d.get("tweet"):
                return d
        except Exception:
            pass
        time.sleep(3 + t * 3)  # fxtwitter は断続的に 404 / 空を返すのでリトライ
    return d


def crawl(uid, ep, extra):
    """プロフィールの一覧をカーソルで最後まで辿る。
    fxtwitter は途中で 404（cursor も null）を断続的に返すので、404 は終端と区別して
    同じカーソルで取り直す（code 200 で結果が空のときだけ終端）。"""
    allr, cur = {}, None
    for _ in range(400):
        u = f"{API}/2/profile/{uid}/{ep}?{extra}" + ("&cursor=" + urllib.parse.quote(cur) if cur else "")
        d, nc, new = {}, None, 0
        for attempt in range(10):
            o = subprocess.run(["curl", "-s", "-m", "30", u], capture_output=True, text=True).stdout
            try:
                d = json.loads(o)
            except Exception:
                d = {}
            if d.get("code") == 200:
                break
            time.sleep(2 + attempt * 2)
        nc = (d.get("cursor") or {}).get("bottom")
        for r in d.get("results", []):
            if r["id"] not in allr:
                allr[r["id"]] = r
                new += 1
        print(ep, extra, len(allr), flush=True)
        if d.get("code") != 200:
            print("警告: 取得が途中で失敗", ep, extra, file=sys.stderr)
            break
        if not d.get("results") or not nc or nc == cur:
            break
        cur = nc
    return allr


def kind(r):
    # replying_to が空で返る返信があるので、本文の先頭 @ / 表示範囲の開始位置でも判定する
    rng = (r.get("raw_text") or {}).get("display_text_range", [0])[0]
    if r.get("replying_to") or r["text"].startswith("@") or rng > 0:
        return "リプライ"
    if r.get("quote"):
        return "引用"
    return "ポスト"


def main():
    uid = sys.argv[1].rstrip("/").split("/")[-1].split("?")[0].lstrip("@")
    out = sys.argv[2]
    os.makedirs(out, exist_ok=True)
    prof = get(f"{API}/{uid}")["user"]
    name = re.sub(r"[^\w぀-ヿ一-鿿ー]", "", prof["name"]) or uid  # 絵文字などを除く
    print(prof["name"], prof["statuses"] if "statuses" in prof else prof.get("tweets"), "件")

    m = {}
    # 通常タイムライン + 返信込み + メディアタブ（取得元ごとに取れる範囲が違うので合算する）
    for ep, ex in (("statuses", ""), ("statuses", "with_replies=true"), ("media", "")):
        m.update(crawl(uid, ep, ex))
    own = [r for r in m.values() if r["author"]["screen_name"].lower() == uid.lower() and not r.get("reposted_by")]
    own.sort(key=lambda r: r["created_timestamp"])

    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XImg
    from openpyxl.styles import Font, Alignment, PatternFill
    from PIL import Image

    imgdir = os.path.join(out, "_img")
    os.makedirs(imgdir, exist_ok=True)

    def thumb(r, i, p):
        f = os.path.join(imgdir, f'{r["id"]}_{i}.jpg')
        if not os.path.exists(f):
            u = p["url"] if p["type"] == "photo" else p.get("thumbnail_url")
            subprocess.run(["curl", "-s", "-o", f, u])
        im = Image.open(f).convert("RGB")
        im.thumbnail((200, 200))
        return im

    wb = Workbook()
    ws = wb.active
    ws.title = f"{name}_{uid}"[:31]
    ws.append(["No.", "投稿日時(JST)", "種別", "本文", "いいね", "リプ数", "リポスト", "閲覧数", "画像", "画像枚数", "画像URL", "投稿URL"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for n, r in enumerate(own, 1):
        ms = (r.get("media") or {}).get("all", [])
        dt = datetime.fromtimestamp(r["created_timestamp"], JST).replace(tzinfo=None)
        urls = "\n".join(p["url"] if p["type"] == "photo" else p["url"] + " (動画)" for p in ms)
        ws.append([n, dt, kind(r), r["text"], r.get("likes"), r.get("replies"), r.get("reposts"), r.get("views"), "", len(ms), urls, r["url"]])
        row = ws.max_row
        ws.cell(row, 2).number_format = "yyyy/mm/dd hh:mm"
        for c in ws[row]:
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (4, 11))
        if ms:
            ws.row_dimensions[row].height = 105
            ths = [thumb(r, i, p) for i, p in enumerate(ms)]
            W = sum(t.width for t in ths) + 4 * (len(ths) - 1)
            H = max(t.height for t in ths)
            cv = Image.new("RGB", (W, H), "white")
            x = 0
            for t in ths:  # 複数枚は横に並べて1枚のサムネにする
                cv.paste(t, (x, 0))
                x += t.width + 4
            f = os.path.join(imgdir, f'{r["id"]}_all.png')
            cv.save(f)
            im = XImg(f)
            s = 100 / H
            im.width, im.height = W * s, H * s
            im.anchor = f"I{row}"
            ws.add_image(im)
    for col, w in zip("ABCDEFGHIJKL", [6, 17, 9, 60, 8, 8, 9, 9, 30, 8, 40, 36]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{ws.max_row}"
    path = os.path.join(out, f"{name}_{uid}_全ポスト.xlsx")
    wb.save(path)
    print(path, len(own), "件 /", sum(1 for r in own if r.get("media")), "件に画像")


main()
